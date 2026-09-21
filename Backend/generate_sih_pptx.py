"""
GhatNetra AI - Automated SIH PowerPoint Generator
SISTec Innovation Hackathon (SIH) - Sagar Institute of Science and Technology, Bhopal
Client / Problem Statement: MPSTDC, Government of Madhya Pradesh
Theme: AI & Intelligent Systems
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette (SISTec Innovation Hackathon & MPSTDC Theme)
    NAVY_DARK = RGBColor(15, 23, 42)      # #0f172a
    BLUE_ACCENT = RGBColor(30, 58, 138)   # #1e3a8a
    SAFFRON_GOLD = RGBColor(245, 158, 11) # #f59e0b
    WHITE = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(148, 163, 184)   # #94a3b8
    CARD_BG = RGBColor(30, 41, 59)        # #1e293b
    CARD_BORDER = RGBColor(51, 65, 85)    # #334155

    blank_slide_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, category_text="SISTEC INNOVATION HACKATHON (SIH) | MPSTDC"):
        # Top banner background
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = NAVY_DARK
        top_bar.line.fill.background()

        # Accent Gold Stripe
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.06))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = SAFFRON_GOLD
        stripe.line.fill.background()

        # Category Text
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.3))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = SAFFRON_GOLD

        # Slide Title
        txBox_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.6))
        tf_title = txBox_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    def add_card(slide, left, top, width, height, title, body_bullets, accent_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), height - Inches(0.3))
        tf = txBox.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = accent_color if accent_color else SAFFRON_GOLD
        p_title.space_after = Pt(6)

        for bullet in body_bullets:
            p = tf.add_paragraph()
            p.text = "• " + bullet
            p.font.size = Pt(10.5)
            p.font.color.rgb = WHITE
            p.space_after = Pt(3)

    # -------------------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE (SISTEC INNOVATION HACKATHON)
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY_DARK
    bg1.line.fill.background()

    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(3.4))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "SISTEC INNOVATION HACKATHON (SIH)"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = SAFFRON_GOLD
    p0.space_after = Pt(6)

    p = tf1.add_paragraph()
    p.text = "🕉️ GhatNetra AI (घाट-नेत्र)"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "AI-Powered River Ghat Crowd Management & Decision Support System"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = SAFFRON_GOLD
    p2.space_after = Pt(12)

    p3 = tf1.add_paragraph()
    p3.text = "Problem Statement By: Madhya Pradesh State Tourism Development Corporation Limited (MPSTDC)\nGovernment of Madhya Pradesh | Theme: AI & Intelligent Systems"
    p3.font.size = Pt(13.5)
    p3.font.color.rgb = GRAY_TEXT

    # Metadata Footer Card
    foot = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.5))
    foot.fill.solid()
    foot.fill.fore_color.rgb = CARD_BG
    foot.line.color.rgb = CARD_BORDER

    ft_box = slide1.shapes.add_textbox(Inches(1.0), Inches(5.25), Inches(11.3), Inches(1.2))
    ft_tf = ft_box.text_frame
    ft_p = ft_tf.paragraphs[0]
    ft_p.text = "SAGAR INSTITUTE OF SCIENCE AND TECHNOLOGY (SISTec), BHOPAL"
    ft_p.font.bold = True
    ft_p.font.size = Pt(13)
    ft_p.font.color.rgb = SAFFRON_GOLD

    ft_p2 = ft_tf.add_paragraph()
    ft_p2.text = "Team Innovation Submission • Working AI Computer Vision & Decision Support Prototype\nFocus: Ujjain Ram Ghat (Shipra River) • Omkareshwar • Maheshwar • Sethani Ghat"
    ft_p2.font.size = Pt(11)
    ft_p2.font.color.rgb = WHITE

    # -------------------------------------------------------------------------
    # SLIDE 2: THE REAL-WORLD PROBLEM STATEMENT
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = NAVY_DARK
    bg2.line.fill.background()
    add_header(slide2, "The Real-World Challenge: River Ghat Crowd Congestion")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "🚨 Mass Pilgrimage Hazards",
             [
                 "Religious mass gatherings (Simhastha Kumbh Mela, Mahakal Sawari, Narmada Jayanti) bring 5-10 lakh pilgrims to river ghats.",
                 "Ancient narrow alleys and steep stone steps create sudden choking points.",
                 "Dangerous stampede and drowning risks during holy dips (Snan)."
             ])

    add_card(slide2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "❌ Manual Limitations",
             [
                 "Police rely on manual eye observation, making response reactive instead of proactive.",
                 "No scientific data on how many pilgrims should be admitted per gate per minute.",
                 "Zero automated warning when ghat capacity crosses critical safety limits.",
                 "Lack of evacuation clearance time calculation during emergencies."
             ])

    add_card(slide2, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "🎯 MPSTDC Scope Mandate",
             [
                 "Continuous crowd occupancy estimation via computer vision.",
                 "Automated recommendations for Gate OPEN/CLOSE and entry rates.",
                 "Dynamic alternate route crowd diversion.",
                 "Physics-based safe evacuation time calculation.",
                 "Pre-event crowd movement simulation for proactive planning."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 3: SYSTEM ARCHITECTURE DIAGRAM (EMBEDDED)
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    bg3 = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = NAVY_DARK
    bg3.line.fill.background()
    add_header(slide3, "System Architecture: End-to-End Operational Pipeline")

    arch_img = "system_architecture.png"
    if os.path.exists(arch_img):
        slide3.shapes.add_picture(arch_img, Inches(0.8), Inches(1.6), width=Inches(11.733), height=Inches(5.3))

    # -------------------------------------------------------------------------
    # SLIDE 4: GHAT SPATIAL LAYOUT & 3-ZONE DECOMPOSITION (EMBEDDED)
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    bg4 = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = NAVY_DARK
    bg4.line.fill.background()
    add_header(slide4, "Spatial Ghat Segmentation: 3-Zone Geometry & Geofencing")

    layout_img = "ghat_zones_layout.png"
    if os.path.exists(layout_img):
        slide4.shapes.add_picture(layout_img, Inches(0.8), Inches(1.6), width=Inches(11.733), height=Inches(5.3))

    # -------------------------------------------------------------------------
    # SLIDE 5: COMPUTER VISION PROTOTYPE OUTPUT (EMBEDDED SNAPSHOT)
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    bg5 = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg5.fill.solid()
    bg5.fill.fore_color.rgb = NAVY_DARK
    bg5.line.fill.background()
    add_header(slide5, "Live Prototype Verification: YOLOv11 & Heads-Up Command Bar")

    snap_img = "test_zone_output.jpg"
    if os.path.exists(snap_img):
        slide5.shapes.add_picture(snap_img, Inches(0.8), Inches(1.6), width=Inches(7.2), height=Inches(4.1))

    # Right side explanations
    add_card(slide5, Inches(8.3), Inches(1.6), Inches(4.2), Inches(5.2),
             "🔍 Visual Analytics Highlights",
             [
                 "Green Bounding Boxes: Ultralytics YOLOv11 nano detects pilgrims in real time at 30+ FPS.",
                 "Red Ground Contact Dots: Exact feet coordinates (cx, y2) mapped to stone steps.",
                 "Translucent Zone Overlays: Dynamic green/yellow/red color mapping via cv2.pointPolygonTest.",
                 "Live Command Center HUD: Displays live headcount, occupancy %, Gate 1/2 status, and alternate diversion route."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 6: DECISION SUPPORT FLOWCHART (EMBEDDED)
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    bg6 = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg6.fill.solid()
    bg6.fill.fore_color.rgb = NAVY_DARK
    bg6.line.fill.background()
    add_header(slide6, "AI Decision Engine: Dynamic Gate Control & Evacuation Physics")

    flow_img = "decision_flowchart.png"
    if os.path.exists(flow_img):
        slide6.shapes.add_picture(flow_img, Inches(0.8), Inches(1.6), width=Inches(11.733), height=Inches(5.3))

    # -------------------------------------------------------------------------
    # SLIDE 7: PRE-EVENT 'WHAT-IF' SIMULATION SANDBOX
    # -------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    bg7 = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg7.fill.solid()
    bg7.fill.fore_color.rgb = NAVY_DARK
    bg7.line.fill.background()
    add_header(slide7, "Pre-Event 'What-If' Simulation Sandbox (For Judges & Police)")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "🎛️ Interactive Parameters",
             [
                 "Officials select major festival scenarios (Simhastha Shahi Snan, Mahakal Sawari).",
                 "Adjust crowd sliders from 30 to 500+ simulated pilgrims.",
                 "Vary exit staircase width (1.5m to 10m) and walking speed to test bottleneck risks."
             ])

    add_card(slide7, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "⚡ Real-Time Decision Logic",
             [
                 "Gate 1 transitions: 🟢 OPEN (120/min) -> 🟡 CONTROLLED (60/min) -> 🔴 CLOSED (Emergency).",
                 "Continuous evacuation clearance time countdown ($T = N / (W \times V \times 60)$).",
                 "Dynamic diversion trigger to Narsingh Ghat before panic starts."
             ])

    add_card(slide7, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "🔊 Multilingual Smart PA",
             [
                 "Automated bilingual voice broadcast in Hindi and English.",
                 "Example: 'कृपया ध्यान दें! राम घाट पर क्षमता पूर्ण हो चुकी है। कृपया नरसिंह घाट की ओर प्रस्थान करें।'",
                 "Clear acoustic guidance eliminates rumor-driven stampedes."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 8: 22-FEATURE COMPREHENSIVE MATRIX
    # -------------------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    bg8 = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg8.fill.solid()
    bg8.fill.fore_color.rgb = NAVY_DARK
    bg8.line.fill.background()
    add_header(slide8, "Complete 22-Feature Alignment With MPSTDC Scope")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "👁️ Vision & Zones (1-8)",
             [
                 "1. YOLOv11 Person Detection",
                 "2. Real-Time Head Counting",
                 "3. Multi-Zone Ghat Decomposition",
                 "4. Zone-Wise People Count",
                 "5. Calibrated Zone Capacities",
                 "6. Dynamic Occupancy Math",
                 "7. Status: Low/Mod/High/Critical",
                 "8. Automated Operator Visual Alerts"
             ])

    add_card(slide8, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "🚪 Flow & Forecasting (9-16)",
             [
                 "9. Entry/Exit Flow Management",
                 "10. Simulated Gate OPEN/CLOSE",
                 "11. Automated Entry Restriction",
                 "12. Dynamic Gate Re-Opening",
                 "13. Movement Vectors & Speed",
                 "14. 2D Gaussian Density Heatmaps",
                 "15. Observed Inflow Trend Graphs",
                 "16. 60-Minute Predictive Rush Forecast"
             ])

    add_card(slide8, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "🛡️ Safety & Command (17-22)",
             [
                 "17. AI Security Deployment Directives",
                 "18. Dynamic Crowd Diversion Routes",
                 "19. AI Lost Child / Person Match Demo",
                 "20. Ujjain Riverfront GIS Map",
                 "21. Streamlit Web Command Center",
                 "22. Weather & River Discharge Data"
             ])

    # -------------------------------------------------------------------------
    # SLIDE 9: TECHNOLOGY, PRIVACY & ON-GROUND DEPLOYMENT
    # -------------------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    bg9 = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg9.fill.solid()
    bg9.fill.fore_color.rgb = NAVY_DARK
    bg9.line.fill.background()
    add_header(slide9, "Technology Stack & Ground Deployment Strategy")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "💻 Technology Core",
             [
                 "Python 3.12 (High performance & cross-platform).",
                 "Ultralytics YOLOv11 nano (Fast inference on standard CPU/GPU).",
                 "OpenCV (cv2) for polygon testing, HUD drawing, and heatmap blending.",
                 "Streamlit with custom CSS for intuitive web command center."
             ])

    add_card(slide9, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "⚙️ Edge Hardware & Offline",
             [
                 "Runs on local edge computers (NVIDIA Jetson / Local Ghat Control Room Server).",
                 "100% Offline Operational Capability: Works even during cellular network congestion.",
                 "Drop-In CCTV Integration: Video path swaps directly with live RTSP IP camera feed."
             ])

    add_card(slide9, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "🔒 Privacy & Safety Impact",
             [
                 "Zero Facial Storage: General pilgrim faces are processed in-memory without recording.",
                 "Dial 112 & SDRF Integration: One-click dispatch of quick response motorized rescue boats.",
                 "Zero-Stampede Mission across all major MP riverfronts."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 10: CONCLUSION & READY FOR LIVE DEMO
    # -------------------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_slide_layout)
    bg10 = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = NAVY_DARK
    bg10.line.fill.background()

    th_box = slide10.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(4.5))
    th_tf = th_box.text_frame
    th_tf.word_wrap = True

    tp0 = th_tf.paragraphs[0]
    tp0.text = "SISTEC INNOVATION HACKATHON (SIH) 2026"
    tp0.font.size = Pt(15)
    tp0.font.bold = True
    tp0.font.color.rgb = SAFFRON_GOLD
    tp0.alignment = PP_ALIGN.CENTER
    tp0.space_after = Pt(10)

    tp = th_tf.add_paragraph()
    tp.text = "Thank You, Respected Evaluators! 🙏"
    tp.font.size = Pt(38)
    tp.font.bold = True
    tp.font.color.rgb = WHITE
    tp.alignment = PP_ALIGN.CENTER
    tp.space_after = Pt(14)

    tp2 = th_tf.add_paragraph()
    tp2.text = "GhatNetra AI (घाट-नेत्र)\nAI-Powered River Ghat Crowd Management & Decision Support System\nDesigned for MPSTDC, Government of Madhya Pradesh"
    tp2.font.size = Pt(18)
    tp2.font.bold = True
    tp2.font.color.rgb = SAFFRON_GOLD
    tp2.alignment = PP_ALIGN.CENTER
    tp2.space_after = Pt(20)

    tp3 = th_tf.add_paragraph()
    tp3.text = "READY FOR LIVE DEMONSTRATION:\n1. Standalone OpenCV HUD Pipeline (`zone_detection.py`)\n2. Interactive Web Command Center (`streamlit run app.py`)\n3. Pre-Event 'What-If' Simulation Sandbox"
    tp3.font.size = Pt(13.5)
    tp3.font.color.rgb = GRAY_TEXT
    tp3.alignment = PP_ALIGN.CENTER

    output_filename = "GhatNetra_AI_SIH2026_Presentation.pptx"
    prs.save(output_filename)
    print(f"[SUCCESS] SISTec SIH Presentation generated with diagrams: {output_filename}")

if __name__ == "__main__":
    create_presentation()
